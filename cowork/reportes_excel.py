from datetime import date, datetime
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .fechas import hora_local_sin_zona


class GeneradorExcel:
    """Convierte filas ya validadas en un libro Excel."""

    def generar(self, filas):
        libro = Workbook()
        hoja = libro.active
        hoja.title = 'Reservas'
        self._cargar_filas(hoja, filas)
        self._configurar_hoja(hoja)
        libro.create_sheet('Información').append(['Zona horaria', timezone.get_current_timezone_name()])
        salida = BytesIO()
        libro.save(salida)
        return salida.getvalue()

    def _cargar_filas(self, hoja, filas):
        for fila in filas:
            valores = [hora_local_sin_zona(valor) if isinstance(valor, datetime) else valor for valor in fila]
            hoja.append(valores)
            for celda in hoja[hoja.max_row]:
                self._formatear_celda(celda)

    @staticmethod
    def _formatear_celda(celda):
        if isinstance(celda.value, str):
            # El contenido del usuario se conserva como texto, incluso si empieza con '='.
            celda.data_type = 's'
        elif isinstance(celda.value, datetime):
            celda.number_format = 'yyyy-mm-dd hh:mm'
        elif isinstance(celda.value, date):
            celda.number_format = 'yyyy-mm-dd'
        elif isinstance(celda.value, float):
            celda.number_format = '0.00'

    @staticmethod
    def _configurar_hoja(hoja):
        for celda in hoja[1]:
            celda.font = Font(bold=True)
            celda.fill = PatternFill('solid', fgColor='E2E8F0')
        for numero in range(1, hoja.max_column + 1):
            hoja.column_dimensions[get_column_letter(numero)].width = 24
        hoja.freeze_panes = 'A2'
        hoja.auto_filter.ref = hoja.dimensions
