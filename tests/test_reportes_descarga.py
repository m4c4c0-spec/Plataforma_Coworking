from datetime import timedelta
from io import BytesIO

from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from cowork.models import Reserva
from .test_gestion import DatosCoworkTests


class DescargaReportesTests(DatosCoworkTests):
    def setUp(self):
        super().setUp()
        reserva = self.nueva_reserva()
        reserva.confirmar(self.admin)
        self.fecha = timezone.localdate(self.inicio).isoformat()

    def parametros(self, **cambios):
        return {'tipo': 'reservas', 'formato': 'xlsx', 'desde': self.fecha, 'hasta': self.fecha, **cambios}

    def test_pantalla_muestra_formulario_y_descarga_excel(self):
        self.client.force_login(self.admin.cuenta)
        self.assertContains(self.client.get(reverse('reportes')), 'Descargar')
        respuesta = self.client.get(reverse('reportes'), self.parametros())
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn('attachment; filename="reporte-reservas.xlsx"', respuesta['Content-Disposition'])
        hoja = load_workbook(BytesIO(respuesta.content))['Reservas']
        self.assertEqual(hoja.max_row, 2)
        self.assertEqual(hoja['G2'].value, Reserva.Estado.CONFIRMADA.label)

    def test_pantalla_descarga_pdf_diario(self):
        self.client.force_login(self.admin.cuenta)
        respuesta = self.client.get(reverse('reportes'), self.parametros(tipo='diario', formato='pdf'))
        self.assertEqual(respuesta['Content-Type'], 'application/pdf')
        self.assertTrue(respuesta.content.startswith(b'%PDF-'))

    def test_pantalla_muestra_errores_de_validacion(self):
        self.client.force_login(self.admin.cuenta)
        anterior = (self.inicio - timedelta(days=1)).date().isoformat()
        respuesta = self.client.get(reverse('reportes'), self.parametros(hasta=anterior))
        self.assertContains(respuesta, 'La fecha final no puede ser anterior a la inicial.')
        respuesta = self.client.get(reverse('reportes'), self.parametros(tipo='diario', desde=''))
        self.assertContains(respuesta, 'El reporte diario requiere las fechas desde y hasta.')

    def test_miembro_no_accede_a_reportes(self):
        self.client.force_login(self.miembro.cuenta)
        self.assertEqual(self.client.get(reverse('reportes')).status_code, 403)
        api = APIClient(HTTP_HOST='localhost')
        api.force_login(self.miembro.cuenta)
        self.assertEqual(api.get(reverse('api:reporte-list'), self.parametros()).status_code, 403)

    def test_api_descarga_y_filtra(self):
        api = APIClient(HTTP_HOST='localhost')
        api.force_login(self.admin.cuenta)
        respuesta = api.get(reverse('api:reporte-list'), self.parametros(estado='cancelada'))
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(load_workbook(BytesIO(respuesta.content))['Reservas'].max_row, 1)
        respuesta = api.get(reverse('api:reporte-list'), self.parametros(formato='csv'))
        self.assertEqual(respuesta.status_code, 400)
        self.assertIn('formato', respuesta.json())


class RutasRouterTests(DatosCoworkTests):
    def setUp(self):
        self.client = APIClient(HTTP_HOST='localhost')
        self.client.force_login(self.admin.cuenta)

    def test_raiz_de_la_api_lista_los_recursos(self):
        datos = self.client.get(reverse('api:api-root'), HTTP_ACCEPT='application/json').json()
        for recurso in (
            'sedes', 'tipos-espacio', 'usuarios', 'espacios', 'reservas',
            'historial-reservas', 'eventos-auditoria', 'reportes',
        ):
            self.assertIn(recurso, datos)

    def test_ids_no_numericos_no_coinciden(self):
        self.assertEqual(self.client.get('/cowork/espacios/abc/').status_code, 404)

    def test_router_genera_delete_en_todos_los_recursos(self):
        sede = reverse('api:sede-detail', args=[self.sede.pk])
        self.assertEqual(self.client.delete(sede).status_code, 400)
        self.assertEqual(self.client.delete(reverse('api:espacio-detail', args=[self.espacio.pk])).status_code, 204)
        self.assertEqual(self.client.delete(sede).status_code, 204)
        self.assertEqual(self.client.delete(reverse('api:usuario-detail', args=[self.otro.pk])).status_code, 204)
