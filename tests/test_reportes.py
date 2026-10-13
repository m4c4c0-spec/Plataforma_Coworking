from datetime import date, datetime, timedelta
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.test import SimpleTestCase, TestCase
from django.utils import timezone
from openpyxl import load_workbook

from cowork.exportadores import ExportadorReportes, ReporteExportar
from cowork.models import Administrador, Espacio, EventoAuditoria, Reserva, Sede, TipoEspacio, Usuario
from cowork.reportes_dto import DatosReporte, FiltroReporteDTO, Formato, RangoFecha
from cowork.services import ServicioReportes


class ContratosReporteTests(SimpleTestCase):
    def test_filtro_rechaza_identificadores_estados_y_periodos_invalidos(self):
        filtros = (
            FiltroReporteDTO(usuario_id=True), FiltroReporteDTO(espacio_id=0),
            FiltroReporteDTO(estado='desconocido'),
            FiltroReporteDTO(desde=date(2026, 10, 9), hasta=date(2026, 10, 8)),
            FiltroReporteDTO(desde='2026-10-08'),
        )
        for filtro in filtros:
            with self.subTest(filtro=filtro), self.assertRaises(ValidationError):
                filtro.como_filtros()

    def test_fechas_sin_hora_y_filtro_abierto_usan_zona_del_proyecto(self):
        with timezone.override('America/Santiago'):
            filtros = FiltroReporteDTO(desde=date(2026, 10, 8)).como_filtros()
            self.assertEqual(timezone.localtime(filtros['desde']).hour, 0)
            self.assertNotIn('hasta', filtros)
            rango = RangoFecha(datetime(2026, 10, 8, 10), datetime(2026, 10, 8, 11))
            self.assertTrue(timezone.is_aware(rango.desde))
            self.assertEqual(len(list(rango.dias())), 1)

    def test_exportador_rechaza_formato_y_filas_invalidas(self):
        exportador = ExportadorReportes()
        with self.assertRaises(TypeError):
            ReporteExportar()
        with self.assertRaises(ValidationError):
            exportador.exportar([], 'csv')
        with self.assertRaises(ValidationError):
            exportador.exportar([{'password': 'dato ajeno'}], Formato.EXCEL)

    def test_exportaciones_vacias_siguen_siendo_documentos_validos(self):
        exportador = ExportadorReportes()
        pdf = exportador.generar_pdf([])
        self.assertTrue(pdf.startswith(b'%PDF-'))
        self.assertTrue(pdf.rstrip().endswith(b'%%EOF'))
        libro = load_workbook(BytesIO(exportador.generar_excel([])))
        self.assertEqual(tuple(celda.value for celda in libro['Reservas'][1]), DatosReporte.columnas)
        self.assertEqual(libro['Reservas'].max_row, 1)
        libro.close()


class ReportesAdministracionTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        def usuario(nombre, rol=Usuario.Rol.MIEMBRO):
            cuenta = get_user_model().objects.create_user(username=nombre)
            return Usuario.objects.create(
                cuenta=cuenta, nombre=nombre, correo=f'{nombre}@example.com', rol=rol,
            )

        cls.actor = usuario('admin_reportes', Usuario.Rol.ADMINISTRADOR)
        cls.miembro = usuario('miembro_reportes')
        cls.otro = usuario('otro_reportes')
        cls.administrador = Administrador.objects.get(pk=cls.actor.pk)
        sede = Sede.objects.create(nombre='Centro', direccion='Calle 1', ciudad='Santiago')
        tipo = TipoEspacio.objects.create(nombre='Sala de reportes')
        cls.espacio = Espacio.objects.create(
            sede=sede, tipo=tipo, nombre='Sala principal', capacidad=5,
            observaciones_privadas='Dato privado que no debe exportarse',
        )
        cls.inicio = timezone.make_aware(datetime(2026, 10, 8, 23))
        cls.fin = cls.inicio + timedelta(hours=3)

    def crear_reserva(self, **cambios):
        datos = dict(usuario=self.miembro, espacio=self.espacio, inicio=self.inicio, fin=self.fin)
        datos.update(cambios)
        reserva = Reserva(**datos)
        reserva.guardar(datos['usuario'])
        return reserva

    def test_reporte_filtra_relaciones_estado_y_reservas_que_cruzan_el_periodo(self):
        reserva = self.crear_reserva()
        self.crear_reserva(usuario=self.otro, inicio=self.fin, fin=self.fin + timedelta(hours=1))
        datos = ServicioReportes.generacionreporte_reserva_espacio(FiltroReporteDTO(
            usuario_id=self.miembro.pk, espacio_id=self.espacio.pk,
            estado=Reserva.Estado.PENDIENTE,
            desde=self.inicio + timedelta(minutes=30), hasta=self.fin,
        ), self.actor)
        self.assertEqual(len(datos), 1)
        self.assertEqual(datos[0].id_reserva, reserva.pk)
        self.assertEqual(datos[0].inicio, self.inicio)
        self.assertNotIn(self.espacio.observaciones_privadas, str(datos))
        self.assertEqual(EventoAuditoria.objects.count(), 2)

    def test_reportes_revalidan_el_rol_y_la_actividad_persistidos(self):
        periodo = RangoFecha(self.inicio, self.fin)
        self.miembro.rol = Usuario.Rol.ADMINISTRADOR
        for metodo, argumento in (
            (ServicioReportes.generacionreporte_reserva_espacio, FiltroReporteDTO()),
            (ServicioReportes.generacionreportereserva_espacio_diarias, periodo),
        ):
            with self.subTest(metodo=metodo), self.assertRaises(PermissionDenied):
                metodo(argumento, self.miembro)
        Usuario.objects.filter(pk=self.actor.pk).update(activo=False)
        with self.assertRaises(PermissionDenied):
            ServicioReportes.resumirReservasDiarias(periodo, self.actor)

    def test_diario_reparte_reserva_entre_dias_y_recorta_extremos(self):
        reserva = self.crear_reserva()
        reserva.confirmar(self.actor)
        self.crear_reserva(inicio=self.fin, fin=self.fin + timedelta(hours=1))
        periodo = RangoFecha(self.inicio + timedelta(minutes=30), self.fin - timedelta(minutes=30))
        datos = ServicioReportes.resumirReservasDiarias(periodo, self.actor)
        self.assertEqual([(fila.fecha, fila.reservas, fila.horas) for fila in datos], [
            (date(2026, 10, 8), 1, 0.5), (date(2026, 10, 9), 1, 1.5),
        ])
        # El límite hasta es exclusivo: una medianoche no agrega un día extra.
        medianoche = timezone.make_aware(datetime(2026, 10, 9))
        datos = ServicioReportes.resumirReservasDiarias(RangoFecha(self.inicio, medianoche), self.actor)
        self.assertEqual(len(datos), 1)
        self.assertEqual(datos[0].horas, 1)

    def test_diario_excluye_canceladas_y_pendientes_e_incluye_espacios_sin_uso(self):
        reserva = self.crear_reserva()
        reserva.cancelar(self.miembro, 'Cambio de planes')
        self.crear_reserva()
        datos = ServicioReportes.resumirReservasDiarias(RangoFecha(self.inicio, self.fin), self.actor)
        self.assertTrue(all(fila.reservas == 0 and fila.horas == 0 for fila in datos))

    def test_diario_incluye_finalizadas(self):
        reserva = self.crear_reserva()
        reserva.confirmar(self.actor)
        with patch('cowork.models.timezone.now', return_value=self.fin + timedelta(minutes=1)):
            reserva.finalizar(self.actor)
        datos = ServicioReportes.resumirReservasDiarias(RangoFecha(self.inicio, self.fin), self.actor)
        self.assertEqual(sum(fila.horas for fila in datos), 3)

    def test_cambio_horario_de_santiago_suma_horas_transcurridas(self):
        with timezone.override('America/Santiago'):
            periodo = RangoFecha(date(2026, 9, 6), date(2026, 9, 7))
            reserva = self.crear_reserva(inicio=periodo.desde, fin=periodo.hasta)
            reserva.confirmar(self.actor)
            datos = ServicioReportes.resumirReservasDiarias(periodo, self.actor)
            self.assertEqual(len(datos), 1)
            self.assertEqual(datos[0].horas, 23)

    def test_excel_se_puede_leer_y_conserva_nombres_como_texto(self):
        self.miembro.nombre = '=1+1'
        self.miembro.save(update_fields=['nombre'])
        self.crear_reserva()
        datos = ServicioReportes.generacionreporte_reserva_espacio(FiltroReporteDTO(), self.actor)
        libro = load_workbook(BytesIO(ExportadorReportes().exportar(datos, Formato.EXCEL)))
        hoja = libro['Reservas']
        self.assertEqual(hoja['B2'].value, '=1+1')
        self.assertEqual(hoja['B2'].data_type, 's')
        self.assertEqual(hoja['E2'].value, timezone.localtime(self.inicio).replace(tzinfo=None))
        self.assertEqual(libro['Información']['B1'].value, timezone.get_current_timezone_name())
        libro.close()

    def test_diario_exporta_pdf_y_excel_con_los_totales_correctos(self):
        reserva = self.crear_reserva()
        reserva.confirmar(self.actor)
        periodo = RangoFecha(self.inicio, self.fin)
        pdf = ServicioReportes.generacionreportereserva_espacio_diarias(periodo, self.actor)
        self.assertTrue(pdf.startswith(b'%PDF-'))
        contenido = ServicioReportes.generacionreportereserva_espacio_diarias(periodo, self.actor, Formato.EXCEL)
        libro = load_workbook(BytesIO(contenido))
        self.assertEqual(libro['Reservas']['E2'].value, 1)
        self.assertEqual(libro['Reservas']['E3'].value, 2)
        libro.close()

    def test_proxy_comparte_identidad_y_filtra_solo_administradores(self):
        self.assertEqual(self.administrador.pk, self.actor.pk)
        self.assertEqual(self.administrador.cuenta_id, self.actor.cuenta_id)
        self.assertEqual(Administrador.objects.count(), 1)
        self.assertFalse(Administrador.objects.filter(pk=self.miembro.pk).exists())
        cuenta = get_user_model().objects.create_user(username='nuevo_admin_reportes')
        nuevo = Administrador(cuenta=cuenta, nombre='Nuevo', correo='nuevo_admin@example.com')
        nuevo.guardar(self.actor)
        self.assertEqual(Usuario.objects.get(pk=nuevo.pk).rol, Usuario.Rol.ADMINISTRADOR)

    def test_gestion_horario_reprograma_pendiente_y_registra_cambio(self):
        reserva = self.crear_reserva()
        actual = self.administrador.gestion_horario(
            self.fin, self.fin + timedelta(hours=1), reserva=reserva,
        )
        reserva.refresh_from_db()
        self.assertEqual(actual.pk, reserva.pk)
        self.assertEqual(reserva.inicio, self.fin)
        self.assertEqual(reserva.fin, self.fin + timedelta(hours=1))
        self.assertEqual(reserva.historial.count(), 1)
        evento = EventoAuditoria.objects.get(accion='actualizar')
        self.assertEqual(evento.actor_id, self.actor.pk)
        self.assertNotEqual(evento.valores_anteriores['inicio'], evento.valores_nuevos['inicio'])

    def test_gestion_horario_rechaza_solapamientos_y_reservas_confirmadas(self):
        primera = self.crear_reserva()
        segunda = self.crear_reserva(inicio=self.fin, fin=self.fin + timedelta(hours=2))
        with self.assertRaises(ValidationError):
            self.administrador.gestion_horario(segunda.inicio, segunda.fin, reserva=primera)
        primera.refresh_from_db()
        self.assertEqual(primera.inicio, self.inicio)
        primera.confirmar(self.actor)
        with self.assertRaises(ValidationError):
            self.administrador.gestion_horario(self.inicio, self.fin, reserva=primera)

    def test_gestion_horario_respeta_la_politica_vigente_de_edicion(self):
        reserva = self.crear_reserva()
        original = Usuario.puedeRealizar

        def politica(usuario, accion):
            if usuario.rol == Usuario.Rol.ADMINISTRADOR and accion == 'editar_reserva':
                return False
            return original(usuario, accion)

        with patch.object(Usuario, 'puedeRealizar', politica), self.assertRaises(PermissionDenied):
            self.administrador.gestion_horario(self.fin, self.fin + timedelta(hours=1), reserva=reserva)
        reserva.refresh_from_db()
        self.assertEqual(reserva.inicio, self.inicio)
        self.assertEqual(EventoAuditoria.objects.count(), 1)

    def test_gestion_espacio_registra_actor_y_rechaza_rol_desactualizado(self):
        pk = self.administrador.gestion_espacios(False, espacio=self.espacio)
        self.assertEqual(pk, self.espacio.pk)
        self.espacio.refresh_from_db()
        self.assertEqual(self.espacio.estado, Espacio.Estado.INACTIVO)
        self.assertEqual(EventoAuditoria.objects.get().actor_id, self.actor.pk)
        self.administrador.gestion_espacios(True, espacio=self.espacio)
        Usuario.objects.filter(pk=self.actor.pk).update(rol=Usuario.Rol.MIEMBRO)
        with self.assertRaises(PermissionDenied):
            self.administrador.gestion_espacios(False, espacio=self.espacio)

    def test_gestion_espacio_revierte_estado_si_falla_la_auditoria(self):
        with patch.object(EventoAuditoria, 'registrar', side_effect=RuntimeError('Fallo')):
            with self.assertRaises(RuntimeError):
                self.administrador.gestion_espacios(False, espacio=self.espacio)
        self.espacio.refresh_from_db()
        self.assertEqual(self.espacio.estado, Espacio.Estado.DISPONIBLE)

    def test_varios_administradores_conservan_relacion_por_historial(self):
        cuenta = get_user_model().objects.create_user(username='segundo_admin_reportes')
        segundo = Administrador.objects.create(cuenta=cuenta, nombre='Segundo', correo='segundo@example.com')
        reserva = self.crear_reserva()
        reserva.confirmar(self.administrador)
        reserva.cancelar(segundo, 'Cancelación administrativa')
        self.assertEqual(list(self.administrador.reservas_administradas), [reserva])
        self.assertEqual(list(segundo.reservas_administradas), [reserva])
